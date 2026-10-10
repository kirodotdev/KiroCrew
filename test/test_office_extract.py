"""The deadline-bounded .docx / .pptx extractor child."""

from __future__ import annotations

import io
import json
import sys
import time
import zipfile

import pytest

from kiro_crew import office_extract, office_extract_child

pytestmark = pytest.mark.skipif(sys.platform == "win32", reason="POSIX rlimits")

_FAR = 30.0
_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _docx(body: str) -> bytes:
    xml = f'<?xml version="1.0"?><w:document {_W}><w:body>{body}</w:body></w:document>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", xml)
    return buf.getvalue()


def _soon(secs: float = _FAR) -> float:
    return time.monotonic() + secs


def test_a_real_docx_is_extracted_in_the_child():
    import docx

    doc = docx.Document()
    doc.add_paragraph("Quarterly plan")
    doc.add_paragraph("Second line")
    buf = io.BytesIO()
    doc.save(buf)

    outcome = office_extract.extract_office_text(
        buf.getvalue(), "docx", max_chars=1000, deadline=_soon()
    )

    assert outcome.failure is None
    assert outcome.text == "Quarterly plan\nSecond line"


def test_text_is_capped_at_max_chars():
    body = "".join(f"<w:p><w:r><w:t>{'x' * 50}</w:t></w:r></w:p>" for _ in range(10))

    outcome = office_extract.extract_office_text(
        _docx(body), "docx", max_chars=30, deadline=_soon()
    )

    assert outcome.failure is None
    assert len(outcome.text) == 30


def test_a_parse_that_outruns_the_deadline_is_killed():
    """Nested empty paragraphs make doc_parser's walk quadratic; the deadline
    ends the child instead of waiting for it."""
    n = 40_000
    bomb = _docx("<w:p>" * n + "</w:p>" * n)

    outcome = office_extract.extract_office_text(bomb, "docx", max_chars=1000, deadline=_soon(1.0))

    assert outcome == office_extract.OfficeExtraction("", "timeout")


def test_bytes_that_are_not_a_zip_are_refused_without_a_spawn(monkeypatch):
    monkeypatch.setattr(office_extract, "popen_limited", lambda *a, **k: pytest.fail("spawned"))

    outcome = office_extract.extract_office_text(
        b"%PDF-1.4", "docx", max_chars=10, deadline=_soon()
    )

    assert outcome.failure == "parse"


def test_a_passed_deadline_spawns_nothing(monkeypatch):
    monkeypatch.setattr(office_extract, "popen_limited", lambda *a, **k: pytest.fail("spawned"))

    outcome = office_extract.extract_office_text(
        b"PK\x03\x04", "docx", max_chars=10, deadline=time.monotonic() - 1
    )

    assert outcome.failure == "timeout"


def test_an_unknown_format_is_a_caller_error():
    with pytest.raises(ValueError):
        office_extract.extract_office_text(b"PK\x03\x04", "xlsx", max_chars=10, deadline=_soon())


@pytest.mark.parametrize(
    ("out", "rc", "failure"),
    [
        (b'{"error": "parse", "detail": "BadZipFile"}\n', 3, "parse"),
        (b'{"error": "memory", "detail": "rss"}\n', 3, "memory"),
        (b'{"error": "other"}\n', 3, "protocol"),
        (b"not json", 1, "protocol"),
        (b'{"text": "ok"}\n', 1, "protocol"),
        (b"", -9, "killed"),
    ],
)
def test_decode_maps_each_child_outcome(out, rc, failure):
    assert office_extract._decode(out, b"", rc, max_chars=100).failure == failure


def test_decode_refuses_text_past_the_cap():
    out = json.dumps({"text": "x" * 11}).encode()

    assert office_extract._decode(out, b"", 0, max_chars=10).failure == "protocol"


def test_child_rejects_unknown_arguments():
    with pytest.raises(SystemExit):
        office_extract_child._parse_args(["--format=xlsx", "--max-chars=1", "--max-rss=1"])
    with pytest.raises(SystemExit):
        office_extract_child._parse_args(["--format=docx", "--max-chars=0", "--max-rss=1"])


@pytest.mark.parametrize(
    "argv",
    [
        ["--format", "--max-chars=1", "--max-rss=1"],  # no '=' separator
        ["--bogus=docx", "--max-chars=1", "--max-rss=1"],  # unknown name
        ["--format=docx", "--max-rss=1"],  # missing --max-chars
        ["--format=docx", "--max-chars=nope", "--max-rss=1"],  # non-integer
    ],
)
def test_parse_args_rejects_malformed_input(argv):
    with pytest.raises(SystemExit):
        office_extract_child._parse_args(argv)


def test_parse_args_accepts_a_well_formed_line():
    assert office_extract_child._parse_args(
        ["--format=pptx", "--max-chars=50", "--max-rss=1024"]
    ) == ("pptx", 50, 1024)


def test_emit_writes_one_json_line(capsys):
    office_extract_child._emit({"text": "hello"})

    out = capsys.readouterr().out
    assert out.endswith("\n")
    assert json.loads(out) == {"text": "hello"}


def _run_main(monkeypatch, extract, data=b"doc-bytes", max_chars=1000):
    """Drive child main() in-process: stub the watchdog, feed stdin, patch the parse."""
    import kiro_crew.doc_parser as doc_parser

    monkeypatch.setattr(office_extract_child, "start_rss_watchdog", lambda *_a, **_k: None)
    monkeypatch.setattr(office_extract_child.sys, "stdin", io.TextIOWrapper(io.BytesIO(data)))
    monkeypatch.setattr(doc_parser, "extract_text", extract)
    return office_extract_child.main(
        ["--format=docx", f"--max-chars={max_chars}", "--max-rss=4096"]
    )


def test_main_emits_text_on_success(monkeypatch, capsys):
    seen: list[bytes] = []

    def _extract(name, *, filename, max_chars, fileobj):
        seen.append(fileobj.read())
        return "Extracted body"

    rc = _run_main(monkeypatch, _extract, data=b"PK\x03\x04body")

    assert rc == 0
    assert seen == [b"PK\x03\x04body"]
    assert json.loads(capsys.readouterr().out) == {"text": "Extracted body"}


def test_main_caps_emitted_text_at_max_chars(monkeypatch, capsys):
    rc = _run_main(monkeypatch, lambda *a, **k: "y" * 100, max_chars=10)

    assert rc == 0
    assert json.loads(capsys.readouterr().out) == {"text": "y" * 10}


def test_main_reports_memory_failure(monkeypatch, capsys):
    def boom(*_a, **_k):
        raise MemoryError

    rc = _run_main(monkeypatch, boom)

    assert rc == office_extract_child.EXIT_FAILED
    assert json.loads(capsys.readouterr().out) == {"error": "memory", "detail": "MemoryError"}


def test_main_reports_parse_failure(monkeypatch, capsys):
    def boom(*_a, **_k):
        raise ValueError("bad doc")

    rc = _run_main(monkeypatch, boom)

    assert rc == office_extract_child.EXIT_FAILED
    assert json.loads(capsys.readouterr().out) == {"error": "parse", "detail": "ValueError"}


def test_a_non_positive_cap_is_a_caller_error():
    with pytest.raises(ValueError):
        office_extract.extract_office_text(b"PK\x03\x04", "docx", max_chars=0, deadline=_soon())


def test_a_spawn_failure_is_reported(monkeypatch):
    def _refuse(*_a, **_k):
        raise OSError("no fork")

    monkeypatch.setattr(office_extract, "popen_limited", _refuse)

    outcome = office_extract.extract_office_text(
        b"PK\x03\x04", "docx", max_chars=10, deadline=_soon()
    )

    assert outcome == office_extract.OfficeExtraction("", "spawn")


@pytest.mark.parametrize(
    ("out", "rc", "failure"),
    [
        (b"x" * (10 * 12 + 65), 0, "protocol"),  # output past the read ceiling
        (b"", -24, "cpu"),  # SIGXCPU from RLIMIT_CPU
        (b"", -15, "signal:15"),
        (b"[1, 2]", 0, "protocol"),  # valid JSON, not an object
    ],
)
def test_decode_maps_resource_and_shape_failures(out, rc, failure):
    assert office_extract._decode(out, b"", rc, max_chars=10).failure == failure
